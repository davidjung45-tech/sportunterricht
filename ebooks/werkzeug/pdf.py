"""Baut alle E-Books der Reihe „Sportunterricht“ als PDF (Vollversion + Leseprobe + Cover-Bild).

    python3 werkzeug/pdf.py            # Entwurf (Fußzeile „Entwurf“)
    python3 werkzeug/pdf.py --final    # Verkaufsversion

Texte der Spiele kommen aus der Stammdaten-Excel (zentrale Quelle), Aufbau/Kapitel aus <buch>/inhalt.md.
"""
import base64
import html
import io
import json
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

import markdown
import qrcode
from openpyxl import load_workbook
from qrcode.image.svg import SvgPathImage
from weasyprint import HTML

HIER = Path(__file__).resolve().parent
EBOOKS = HIER.parent
sys.path.insert(0, str(HIER))
from buch import parse, stufen_text  # noqa: E402

APP = EBOOKS.parent / 'sportunterricht-app-v2'
EXCEL = APP / 'data' / 'Sportunterricht_Stammdaten.xlsx'
FONTS = APP / 'node_modules' / '@fontsource'
EIGENE_FONTS = APP / 'src' / 'fonts'  # Atkinson Hyperlegible Next mit ungestrichener Null
# Ohne App-Ordner: Schriften aus website/statisch/fonts (dort liegen alle vier Dateien samt OFL-Lizenz)
WEB_FONTS = EBOOKS.parent / 'website' / 'statisch' / 'fonts'
if not (EIGENE_FONTS / 'atkinson-400.woff2').exists():
    EIGENE_FONTS = WEB_FONTS
OUT = EBOOKS / 'ausgabe'
FINAL = '--final' in sys.argv
KANAL = 'https://www.youtube.com/@Sportunterricht2022'
EINSTELLUNGEN = EBOOKS.parent / 'website' / 'einstellungen.json'
WEBSITE = (json.loads(EINSTELLUNGEN.read_text()).get('domain') or '').rstrip('/') or None if EINSTELLUNGEN.exists() else None
STAND = date.today().strftime('%m/%Y')

REIHE = [
    {'slug': 'spielesammlung', 'band': 1, 'kurz': '50 Spiele', 'titel': '50 kreative & kooperative Spiele mit wenig Material',
     'unter': 'Die meistgesehenen Spiele des Kanals – mit Video-Demo, Ansagetext und Differenzierung', 'farbe': '#1f5bd8', 'preis': '14,90 €',
     'einheit': 'Spiele', 'text': 'Aufwärmen, Kooperation, große Gruppen, kleine Hallen, Abschluss – die Klassiker des Kanals, ausgewählt nach echten Aufrufzahlen.'},
    {'slug': 'aufwaermen', 'band': 2, 'kurz': 'Aufwärmen', 'titel': 'Aufwärmen in 5 Minuten',
     'unter': '40 Fang-, Lauf- und Reaktionsspiele, die sofort funktionieren', 'farbe': '#c8402a', 'preis': '9,90 €',
     'einheit': 'Spiele', 'text': 'Schnell erklärt, sofort in Bewegung: Fangspiele, Staffeln, Reaktionsspiele und Ideen für kleine Hallen.'},
    {'slug': 'turnen', 'band': 3, 'kurz': 'Turnen', 'titel': 'Turnen sicher unterrichten',
     'unter': 'Methodische Reihen, Hilfestellung und Stationen für Boden, Kasten, Reck, Barren und Minitramp', 'farbe': '#5b3fa8', 'preis': '14,90 €',
     'einheit': 'Übungen', 'text': 'Vom Päckchen zur Rolle, von der Vorübung zum Hüftaufschwung: jede Übung mit methodischer Reihe, Hilfestellung und typischen Fehlern.'},
    {'slug': 'fitness', 'band': 4, 'kurz': 'Fitness', 'titel': 'Fit im Schulsport',
     'unter': 'Kraft, Ausdauer und Koordination mit Spaß – Zirkel, Staffeln, Challenges', 'farbe': '#b8531a', 'preis': '12,90 €',
     'einheit': 'Übungen', 'text': 'Kraft, Ausdauer, Koordination und Leichtathletik kindgerecht – mit sechs fertigen Zirkeln für die Stunde.'},
    {'slug': 'praktikum', 'band': 5, 'kurz': 'Praktikums-Kit', 'titel': 'Das Praktikums-Kit Bewegung und Sport',
     'unter': '12 fertige Stundenbilder für Volksschule und Sekundarstufe I – plus Leitfaden für Planung und Reflexion', 'farbe': '#167a4b', 'preis': '19,90 €',
     'einheit': 'Stundenbilder', 'text': 'Für Praktikum und Berufseinstieg: 12 Stundenbilder mit Lernzielen, Organisation, Differenzierung, Sicherheit und Lehrplanbezug.'},
    {'slug': 'alle-dabei', 'band': 6, 'kurz': 'Alle dabei', 'titel': 'Alle dabei, alle sicher',
     'unter': 'Sicherheit, Angst, Aggression, Motivation, Heterogenität, Nicht-Aktive, Koedukation, „Strafen“ und Rituale im Sportunterricht – mit Formulierungen, Fallbeispielen und Kopiervorlagen', 'farbe': '#0f6d7a', 'preis': '24,90 €',
     'einheit': 'Kapitel', 'text': 'Die schwierigen Situationen im Turnsaal: neun Kapitel mit Strategien, wörtlichen Formulierungen, Fallbeispielen, Checklisten und Kopiervorlagen.'},
]


def _euro(s):
    return float(s.replace(' €', '').replace(',', '.'))


def _fmt(x):
    return f'{x:.2f}'.replace('.', ',') + ' €'


KOMPLETT = {'slug': 'komplettpaket', 'titel': 'Das Komplettpaket', 'preis': '59,90 €', 'statt': _fmt(sum(_euro(b['preis']) for b in REIHE))}
ZAHLWORT = {5: 'fünf', 6: 'sechs', 7: 'sieben', 8: 'acht'}
N_BAENDE = len(REIHE)
BAENDE_WORT = ZAHLWORT.get(N_BAENDE, str(N_BAENDE))


def font(fam, datei):
    quelle = FONTS / fam / 'files' / datei
    if not quelle.exists():  # z. B. 'barlow-condensed-latin-700-normal.woff2' → website/statisch/fonts/barlow-700.woff2
        quelle = WEB_FONTS / ('barlow-' + datei.split('-latin-')[1].split('-')[0] + '.woff2')
    return quelle.as_uri()


FONT_CSS = f'''
@font-face {{ font-family: 'Atkinson'; font-weight: 400; src: url('{EIGENE_FONTS.as_uri()}/atkinson-400.woff2') format('woff2'); }}
@font-face {{ font-family: 'Atkinson'; font-weight: 700; src: url('{EIGENE_FONTS.as_uri()}/atkinson-700.woff2') format('woff2'); }}
@font-face {{ font-family: 'Barlow'; font-weight: 600; src: url('{font("barlow-condensed", "barlow-condensed-latin-600-normal.woff2")}') format('woff2'); }}
@font-face {{ font-family: 'Barlow'; font-weight: 700; src: url('{font("barlow-condensed", "barlow-condensed-latin-700-normal.woff2")}') format('woff2'); }}
'''


def css(farbe, fuss):
    return FONT_CSS + f'''
:root {{ --f: {farbe}; --tinte: #13211b; --t2: #46574f; --t3: #66776f; --linie: #d2ddd7; --hell: #f4f8f5; }}
@page {{ size: A4; margin: 12mm 14mm 14mm 14mm;
  @bottom-left {{ content: "{fuss}"; font: 7.5pt 'Atkinson'; color: #8a9c92 }}
  @bottom-right {{ content: counter(page); font: 8pt 'Barlow'; font-weight: 700; color: #8a9c92 }} }}
@page cover {{ margin: 0; @bottom-left {{ content: none }} @bottom-right {{ content: none }} }}
@page quer {{ size: A4 landscape; margin: 9mm 11mm 12mm 11mm; }}
body {{ font-family: 'Atkinson', sans-serif; font-size: 9pt; line-height: 1.36; color: var(--tinte); }}
h1, h2, h3, h4 {{ font-family: 'Barlow', 'Arial Narrow', sans-serif; font-weight: 700; line-height: 1.08; margin: 0; }}
a {{ color: var(--f); text-decoration: none; }}
p {{ margin: 0 0 1.8mm; }}
ul, ol {{ margin: 0 0 1.8mm; padding-left: 5mm; }}
li {{ margin-bottom: 0.5mm; }}
table {{ border-collapse: collapse; width: 100%; font-size: 8.2pt; margin: 2mm 0 4mm; }}
th {{ background: var(--tinte); color: #fff; text-align: left; padding: 1.6mm 2mm; font-weight: 700; }}
td {{ border-bottom: 0.6pt solid var(--linie); padding: 1.4mm 2mm; vertical-align: top; }}
tr {{ page-break-inside: avoid; }}

.cover {{ page: cover; position: relative; height: 297mm; background: var(--f); color: #fff; padding: 30mm 20mm 0; box-sizing: border-box; page-break-after: always; overflow: hidden; }}
.cover svg.feld {{ position: absolute; right: -40mm; bottom: -10mm; width: 190mm; opacity: .28; }}
.cover .reihe {{ font-family: 'Barlow'; font-weight: 700; letter-spacing: .22em; text-transform: uppercase; font-size: 10pt; opacity: .9; }}
.cover .band {{ display: inline-block; margin-top: 4mm; border: 1.2pt solid rgba(255,255,255,.8); padding: 1mm 3mm; font-family: 'Barlow'; font-weight: 700; letter-spacing: .12em; font-size: 9pt; }}
.cover h1 {{ font-size: 46pt; line-height: .98; margin: 16mm 0 6mm; text-transform: uppercase; letter-spacing: .005em; max-width: 165mm; }}
.cover .unter {{ font-size: 13pt; max-width: 140mm; line-height: 1.35; opacity: .95; }}
.cover .zahlen {{ position: absolute; left: 20mm; bottom: 34mm; display: flex; gap: 10mm; }}
.cover .zahlen div {{ border-left: 2.4pt solid #fff; padding-left: 3mm; font-size: 9pt; }}
.cover .zahlen b {{ display: block; font-family: 'Barlow'; font-size: 26pt; line-height: 1; }}
.cover .autor {{ position: absolute; left: 20mm; bottom: 16mm; font-size: 9pt; opacity: .9; }}
.cover .probe {{ position: absolute; right: 20mm; top: 30mm; background: #fff; color: var(--f); font-family: 'Barlow'; font-weight: 700; padding: 2mm 4mm; font-size: 12pt; letter-spacing: .1em; text-transform: uppercase; }}

.impressum {{ page-break-after: always; font-size: 8pt; color: var(--t2); padding-top: 150mm; }}
.impressum h4 {{ font-size: 10pt; color: var(--tinte); margin: 3mm 0 1mm; }}

.inhalt {{ page-break-after: always; }}
.inhalt h1, .text-seite h1 {{ font-size: 24pt; color: var(--f); margin-bottom: 5mm; text-transform: uppercase; }}
.toc {{ list-style: none; padding: 0; margin: 0; }}
.toc li {{ display: flex; gap: 3mm; border-bottom: 0.6pt dotted var(--linie); padding: 1.6mm 0; }}
.toc li.kap {{ font-family: 'Barlow'; font-weight: 700; font-size: 12pt; color: var(--tinte); border-bottom: 1pt solid var(--tinte); margin-top: 3mm; }}
.toc li a {{ color: inherit; flex: 1; }}
.toc li a::after {{ content: target-counter(attr(href), page); float: right; font-family: 'Barlow'; font-weight: 700; color: var(--t3); }}
.ohne-seiten .toc li a::after {{ content: none; }}
.qr-inline {{ width: 13mm; height: 13mm; vertical-align: middle; margin: 0 1mm; }}
li.check {{ list-style: none; margin-left: -5mm; }}
li.check::before {{ content: ''; display: inline-block; width: 2.8mm; height: 2.8mm; border: 0.9pt solid var(--tinte); border-radius: 0.5mm; margin-right: 2mm; vertical-align: -0.5mm; }}
.anh {{ table-layout: fixed; }}
.toc .nr {{ color: var(--f); font-family: 'Barlow'; font-weight: 700; width: 7mm; }}

.text-seite {{ page-break-after: always; }}
.text-seite h2 {{ font-size: 15pt; margin: 5mm 0 2mm; color: var(--tinte); }}
.text-seite h3 {{ font-size: 12pt; margin: 4mm 0 1.5mm; }}
.legende {{ background: var(--hell); border-radius: 2mm; padding: 3mm 4mm; margin-top: 4mm; font-size: 8.6pt; }}

.kapitel {{ page-break-before: always; margin-bottom: 6mm; padding: 9mm 8mm 7mm; background: var(--f); color: #fff; border-radius: 2mm; position: relative; overflow: hidden; }}
.kapitel .knr {{ font-family: 'Barlow'; font-weight: 700; letter-spacing: .2em; text-transform: uppercase; font-size: 9pt; opacity: .85; }}
.kapitel h2 {{ font-size: 26pt; text-transform: uppercase; margin-top: 2mm; }}
.kapitel .kmeta {{ margin-top: 2mm; opacity: .9; }}
.kapitel svg {{ position: absolute; right: -10mm; top: -8mm; width: 70mm; opacity: .25; }}
.ktext {{ margin-bottom: 4mm; }}
.ktext h3 {{ font-size: 13pt; color: var(--f); margin: 4mm 0 1.5mm; }}
.ktext h4 {{ font-size: 11pt; margin: 3mm 0 1mm; }}

.spiel {{ page-break-inside: avoid; border: 0.8pt solid var(--linie); border-top: 3pt solid var(--f); border-radius: 2mm; padding: 3mm 4mm 1mm; margin-bottom: 4mm; }}
.spiel.lang {{ page-break-inside: auto; }}
.spiel.neu {{ page-break-before: always; }}
.kapitel + .spiel.neu, .kapitel + .ktext + .spiel.neu {{ page-break-before: auto; }}
.skopf {{ display: flex; align-items: baseline; gap: 3mm; margin-bottom: 2mm; }}
.snr {{ font-family: 'Barlow'; font-weight: 700; font-size: 18pt; color: var(--f); line-height: 1; }}
.spiel h3 {{ font-size: 15pt; flex: 1; }}
.sreihe {{ display: flex; gap: 4mm; align-items: flex-start; margin-bottom: 2mm; page-break-inside: avoid; }}
.meta {{ flex: 1; display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 1.6mm 4mm; background: var(--hell); padding: 2mm 3mm; border-radius: 1.5mm; font-size: 8.2pt; }}
.meta span {{ display: block; font-family: 'Barlow'; font-weight: 700; font-size: 7.4pt; text-transform: uppercase; letter-spacing: .09em; color: var(--f); }}
.qr {{ width: 18mm; text-align: center; font-size: 6.6pt; color: var(--t3); }}
.qr img {{ width: 16.5mm; height: 16.5mm; display: block; margin: 0 auto 0.6mm; }}
.lernziel {{ font-size: 8.8pt; color: var(--t2); margin-bottom: 2mm; }}
.lernziel b {{ color: var(--tinte); }}
.spiel h4 {{ font-size: 8.4pt; text-transform: uppercase; letter-spacing: .1em; color: var(--f); margin: 1.8mm 0 0.6mm; }}
.ansage {{ background: color-mix(in srgb, var(--f) 9%, #fff); border-radius: 1.5mm; padding: 1.6mm 3mm; font-style: italic; }}
.ansage p:last-child, .sicher p:last-child {{ margin-bottom: 0; }}
.sicher {{ background: #fbf0d2; border-left: 2.4pt solid #b88407; padding: 1.6mm 3mm; margin: 1.5mm 0 2mm; font-size: 8.6pt; }}

.anhang {{ page-break-before: always; }}
.anhang h2 {{ font-size: 22pt; text-transform: uppercase; color: var(--f); margin-bottom: 3mm; }}
.anh td {{ padding: 1mm 2mm; }}
.nw {{ white-space: nowrap; }}
.anhang h3 {{ font-size: 12pt; margin: 5mm 0 1mm; page-break-after: avoid; }}
.reihe-seite {{ page-break-before: always; }}
.reihe-seite h2 {{ font-size: 22pt; text-transform: uppercase; margin-bottom: 4mm; }}
.baende {{ display: grid; grid-template-columns: 1fr 1fr; gap: 4mm; }}
.bandkarte {{ border-radius: 2mm; padding: 4mm; color: #fff; }}
.bandkarte b {{ font-family: 'Barlow'; font-size: 14pt; display: block; line-height: 1.1; margin: 1mm 0 1.5mm; }}
.bandkarte small {{ opacity: .9; display: block; }}
.bandkarte .p {{ font-family: 'Barlow'; font-weight: 700; font-size: 12pt; margin-top: 2mm; }}
.autorbox {{ display: flex; gap: 6mm; align-items: center; margin-top: 8mm; background: var(--hell); padding: 5mm; border-radius: 2mm; }}
.autorbox img {{ width: 26mm; }}

.sb {{ page: quer; page-break-before: always; }}
.sb-kopf {{ display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2pt solid var(--f); padding-bottom: 1.5mm; margin-bottom: 2mm; }}
.sb-kopf h2 {{ font-size: 18pt; }}
.sb-kopf .snr {{ font-size: 26pt; }}
.sb-daten {{ display: grid; grid-template-columns: repeat(5, 1fr); gap: 2mm; font-size: 7.6pt; line-height: 1.3; margin-bottom: 2mm; }}
.sb-daten div {{ background: var(--hell); padding: 1.6mm 2.4mm; border-radius: 1.5mm; }}
.sb-daten span {{ display: block; font-family: 'Barlow'; font-weight: 700; text-transform: uppercase; letter-spacing: .08em; font-size: 7pt; color: var(--f); }}
.sb table {{ font-size: 7pt; line-height: 1.28; margin: 1mm 0 0; }}
.sb th {{ padding: 1.2mm 1.5mm; }}
.sb td {{ padding: 1mm 1.5mm; }}
.sb td small {{ color: var(--t3); }}
.sb td.z {{ white-space: nowrap; font-family: 'Barlow'; font-weight: 700; font-size: 9pt; }}
.sb tr.p-Einstieg td:first-child {{ border-left: 3pt solid #b88407; }}
.sb tr.p-Aufwärmen td:first-child {{ border-left: 3pt solid #c8402a; }}
.sb tr.p-Hauptteil td:first-child {{ border-left: 3pt solid #1f5bd8; }}
.sb tr.p-Abschluss td:first-child {{ border-left: 3pt solid #167a4b; }}
.sb-detail {{ page: quer; page-break-before: always; columns: 3; column-gap: 6mm; font-size: 7.8pt; line-height: 1.32; }}
.sb-detail p {{ margin: 0 0 1.4mm; }}
.sb-detail .blk {{ break-inside: avoid; margin-bottom: 3mm; }}
.sb-detail h4 {{ font-size: 10pt; color: var(--f); margin-bottom: 1mm; }}
.reflexion {{ break-inside: avoid; border: 0.8pt dashed var(--linie); border-radius: 2mm; padding: 3mm; }}
.reflexion p {{ margin: 0 0 6mm; }}
'''


def qr(url):
    img = qrcode.make(url, image_factory=SvgPathImage, box_size=10, border=1)
    b = io.BytesIO()
    img.save(b)
    return 'data:image/svg+xml;base64,' + base64.b64encode(b.getvalue()).decode()


def md(text):
    t = re.sub(r'\s*\((https://www\.youtube\.com/watch\?v=[\w-]+)\)', lambda m_: f' <img class="qr-inline" src="{qr(m_.group(1))}" alt="QR-Code zum Video"/>', text or '')
    h = markdown.markdown(t, extensions=['tables', 'sane_lists'])
    return h.replace('<li>[ ] ', '<li class="check">')


FELD_SVG = '''<svg class="feld" viewBox="0 0 400 260" xmlns="http://www.w3.org/2000/svg" fill="none" stroke="#fff" stroke-width="3">
<rect x="10" y="10" width="380" height="240" rx="6"/><line x1="200" y1="10" x2="200" y2="250"/><circle cx="200" cy="130" r="42"/>
<rect x="10" y="80" width="60" height="100"/><rect x="330" y="80" width="60" height="100"/><circle cx="200" cy="130" r="4" fill="#fff"/></svg>'''
FELD_KLEIN = FELD_SVG.replace(' class="feld"', '')


# ---------------- Stammdaten ----------------
def lade_excel():
    if not EXCEL.exists():  # Ohne Stammdaten-Excel gelten die Texte aus den inhalt.md-Dateien
        print(f'  Hinweis: {EXCEL.name} nicht gefunden – E-Books werden aus den Markdown-Texten gebaut.')
        return {}, {}
    wb = load_workbook(EXCEL, read_only=True, data_only=True)

    def blatt(name):
        rows = list(wb[name].iter_rows(values_only=True))
        kopf = rows[0]
        return {r[0]: dict(zip(kopf, r)) for r in rows[1:] if r[0]}
    return blatt('Videos'), blatt('Anleitungen')


VIDEOS, ANL = lade_excel()


def ja(x):
    return str(x or '').strip().lower() in ('ja', 'x', '1', 'true')


def eintrag_aus_excel(e):
    """Übernimmt Stufe, Dauer, Stundenteil (Blatt Videos) und alle Texte (Blatt Anleitungen)."""
    v, a = VIDEOS.get(e['id']), ANL.get(e['id'])
    m = dict(e['meta'])
    if v:
        if v.get('Stufe_von') and v.get('Stufe_bis'):
            m['Altersstufe'] = stufen_text(int(v['Stufe_von']), int(v['Stufe_bis']))
        if v.get('Spielzeit_min') and v.get('Spielzeit_max'):
            m['Dauer'] = f"{int(v['Spielzeit_min'])}–{int(v['Spielzeit_max'])} Minuten"
        teile = [p for p in ('Aufwärmen', 'Hauptteil', 'Abschluss') if ja(v.get(p))]
        if teile:
            m['Stundenteil'] = ', '.join(teile)
        if v.get('Intensität'):
            m['Intensität'] = v['Intensität']
    name = e['name']
    abschnitte = e['abschnitte']
    if a:
        name = a.get('Spielname') or name
        for k, spalte in [('Gruppengröße', 'Gruppengröße'), ('Material', 'Material_Text'), ('Lernziel', 'Lernziel')]:
            if a.get(spalte):
                m[k] = a[spalte]
        zeilen = lambda t: [z.strip() for z in str(t or '').split('\n') if z.strip()]  # noqa: E731
        art = a.get('Beschreibung_Art') or 'Spielbeschreibung'
        abschnitte = [(art, a.get('Beschreibung') or '')]
        if a.get('Methodische_Reihe'):
            abschnitte.append(('Methodische Reihe', '\n'.join(f'{i}. {z}' for i, z in enumerate(zeilen(a['Methodische_Reihe']), 1))))
        if a.get('Hilfestellung'):
            abschnitte.append(('Hilfestellung', a['Hilfestellung']))
        if a.get('Typische_Fehler'):
            abschnitte.append(('Typische Fehler', '\n'.join(f'- {z}' for z in zeilen(a['Typische_Fehler']))))
        if a.get('So_erklärst_du_es'):
            abschnitte.append(('So erklärst du es', a['So_erklärst_du_es']))
        if a.get('Variationen'):
            abschnitte.append(('Variationen/Differenzierung', '\n'.join(f'- {z}' for z in zeilen(a['Variationen']))))
        if a.get('Sicherheit'):
            abschnitte.append(('Sicherheitshinweis', a['Sicherheit']))
    return {**e, 'name': name, 'meta': m, 'abschnitte': abschnitte}


# ---------------- Bausteine ----------------
def stufen_bereich(eintraege):
    vs = [VIDEOS.get(e['id']) for e in eintraege]
    von = [int(v['Stufe_von']) for v in vs if v and v.get('Stufe_von')]
    bis = [int(v['Stufe_bis']) for v in vs if v and v.get('Stufe_bis')]
    return f'{min(von)}–{max(bis)}' if von and bis else '1–13'


def cover(b, anzahl, probe=False, stufen=None):
    if b['slug'] == 'praktikum':
        zahlen = '<div><b>12</b>Stundenbilder</div><div><b>7</b>Kapitel Leitfaden</div>'
    else:
        zahlen = f'<div><b>{anzahl}</b>{b["einheit"]}</div><div><b>{anzahl}</b>Video-Demos</div>'
    stufen = stufen or '1–8'
    zahlen += f'<div><b>{stufen}</b>Schulstufe</div>'
    return f'''<section class="cover">{FELD_SVG}
<div class="reihe">Sportunterricht · Die Praxis-Reihe</div><div class="band">BAND {b["band"]}</div>
{'<div class="probe">Leseprobe</div>' if probe else ''}
<h1>{html.escape(b["titel"])}</h1><p class="unter">{html.escape(b["unter"])}</p>
<div class="zahlen">{zahlen}</div>
<div class="autor">David Jungreithmayr · AHS-Sportlehrer in Wien · Lehrbeauftragter für Unterrichtspraktische Studien · YouTube: @Sportunterricht2022</div></section>'''


def impressum(b):
    return f'''<section class="impressum">
<h4>{html.escape(b["titel"])}</h4>
<p>Band {b["band"]} der Reihe „Sportunterricht – Die Praxis-Reihe“ · Version 1.0 · Stand {STAND}</p>
<h4>Nutzung</h4>
<p>Dieses E-Book ist für deinen persönlichen Gebrauch und deinen eigenen Unterricht bestimmt. Ausdrucken für die eigene Stunde ist ausdrücklich erlaubt.
Weitergabe, Veröffentlichung oder Upload (auch in Schul-Clouds) sind nicht gestattet. Für das ganze Kollegium gibt es eine Schullizenz.</p>
<h4>Sicherheit</h4>
<p>Alle Spiele und Übungen wurden im Schulunterricht erprobt. Sie ersetzen nicht die Einschätzung der Lehrkraft vor Ort: Passe Regeln, Raum, Material und
Belastung an deine Gruppe an und beachte die geltenden Sicherheitsbestimmungen für den Schulsport sowie deine Aufsichtspflicht. Eine Haftung für Unfälle
oder Schäden bei der Durchführung wird nicht übernommen.</p>
<h4>Videos</h4>
<p>Die QR-Codes führen zu den Demo-Videos auf dem YouTube-Kanal „Sportunterricht“ (@Sportunterricht2022).</p>
<p>© {date.today().year} David Jungreithmayr – Sportunterricht. Alle Rechte vorbehalten.</p>
</section>'''


def toc(buch, eintraege_je_kap, mit_seitenzahlen=True, zusatz=()):
    z = [f'<section class="inhalt{"" if mit_seitenzahlen else " ohne-seiten"}"><h1>Inhalt</h1><ol class="toc">', '<li class="kap"><a href="#einleitung">Bevor du loslegst</a></li>']
    for i, k in enumerate(buch['kapitel']):
        z.append(f'<li class="kap"><a href="#k{i}">{html.escape(k["titel"])}</a></li>')
        for e in eintraege_je_kap[i]:
            z.append(f'<li><span class="nr">{e["nr"]}</span><a href="#e{e["nr"]}">{html.escape(e["name"])}</a></li>')
    for anker, titel in zusatz:
        z.append(f'<li class="kap"><a href="#{anker}">{titel}</a></li>')
    z.append('</ol></section>')
    return '\n'.join(z)


META_ANZEIGE = [('Altersstufe', 'Schulstufe'), ('Gruppengröße', 'Gruppe'), ('Material', 'Material'), ('Dauer', 'Dauer'), ('Stundenteil', 'Stundenteil'), ('Intensität', 'Intensität')]


def ist_lang(e):
    return sum(len(t or '') for _, t in e['abschnitte']) > 1700


def karte(e, neu=False):
    m = e['meta']
    meta = ''.join(f'<div><span>{lab}</span>{html.escape(str(m.get(k, "")))}</div>' for k, lab in META_ANZEIGE if m.get(k))
    url = m.get('Video-Demo', '')
    teile = []
    laenge = 0
    for label, text in e['abschnitte']:
        if not text:
            continue
        laenge += len(text)
        if label == 'Sicherheitshinweis':
            teile.append(f'<div class="sicher"><b>Sicherheit:</b> {md(text)[3:-4] if text.count(chr(10)) == 0 else md(text)}</div>')
        elif label == 'So erklärst du es':
            teile.append(f'<h4>So erklärst du es</h4><div class="ansage">{md(text)}</div>')
        else:
            titel = 'Leichter / schwerer' if label.startswith('Variationen') else label
            teile.append(f'<h4>{titel}</h4>{md(text)}')
    lang = (' lang' if laenge > 1700 else '') + (' neu' if neu else '')
    lz = f'<p class="lernziel"><b>Lernziel:</b> {html.escape(m["Lernziel"])}</p>' if m.get('Lernziel') else ''
    q = f'<div class="qr"><img src="{qr(url)}"/>Video-Demo</div>' if url else ''
    return f'''<article class="spiel{lang}" id="e{e["nr"]}"><div class="skopf"><span class="snr">{e["nr"]}</span><h3>{html.escape(e["name"])}</h3></div>
<div class="sreihe"><div class="meta">{meta}</div>{q}</div>{lz}{"".join(teile)}</article>'''


def kapitel_kopf(i, k, n, farbe):
    teil = k['titel'].split(' – ', 1)
    knr, name = (teil[0], teil[1]) if len(teil) == 2 else ('', k['titel'])
    kmeta = f'<div class="kmeta">{n} Einträge</div>' if n else ''
    return f'<section class="kapitel" id="k{i}">{FELD_KLEIN}<div class="knr">{html.escape(knr)}</div><h2>{html.escape(name)}</h2>{kmeta}</section>'


def anhang(eintraege):
    def zeile(e):
        m = e['meta']
        stufe = m.get('Altersstufe', '').replace(' Schulstufe', '')
        dauer = m.get('Dauer', '').replace(' Minuten', ' min')
        return f'<tr><td>{e["nr"]}</td><td><b>{html.escape(e["name"])}</b></td><td class="nw">{html.escape(stufe)}</td><td class="nw">{html.escape(dauer)}</td><td>{html.escape(m.get("Material", ""))}</td></tr>'
    kopf = '<table class="anh"><thead><tr><th style="width:8mm">Nr.</th><th style="width:52mm">Titel</th><th style="width:18mm">Stufe</th><th style="width:20mm">Dauer</th><th>Material</th></tr></thead><tbody>'
    z = ['<section class="anhang" id="anhang"><h2>Übersicht zum Nachschlagen</h2>']
    for teil in ('Aufwärmen', 'Hauptteil', 'Abschluss'):
        passt = [e for e in eintraege if teil in e['meta'].get('Stundenteil', '')]
        if passt:
            z.append(f'<h3>Geeignet für: {teil} ({len(passt)})</h3>' + kopf + ''.join(zeile(e) for e in passt) + '</tbody></table>')
    ohne = [e for e in eintraege if re.match(r'^\s*keine', e['meta'].get('Material', ''), re.I)]
    if ohne:
        z.append(f'<h3>Ganz ohne Material ({len(ohne)})</h3>' + kopf + ''.join(zeile(e) for e in ohne) + '</tbody></table>')
    z.append('</section>')
    return '\n'.join(z)


def reihe_seite(aktuell):
    karten = []
    for b in REIHE:
        karten.append(f'<div class="bandkarte" style="background:{b["farbe"]}"><small>BAND {b["band"]}{" · dieses Buch" if b["slug"] == aktuell else ""}</small><b>{html.escape(b["titel"])}</b><small>{html.escape(b["text"])}</small><div class="p">{b["preis"]}</div></div>')
    karten.append(f'<div class="bandkarte" style="background:#13211b"><small>ALLE {N_BAENDE} BÄNDE</small><b>Das Komplettpaket</b><small>Alle {BAENDE_WORT} Bände plus Jahresplaner als Bonus – einzeln {KOMPLETT["statt"]}.</small><div class="p">{KOMPLETT["preis"]}</div></div>')
    ziel = WEBSITE or KANAL
    return f'''<section class="reihe-seite" id="reihe"><h2>Die Praxis-Reihe</h2><div class="baende">{"".join(karten)}</div>
<div class="autorbox"><img src="{qr(ziel)}"/><div><h3>Über den Autor</h3><p>David Jungreithmayr unterrichtet Bewegung und Sport an einer AHS in Wien, leitet regelmäßig Wintersportwochen und hält an der Universität Wien Lehrveranstaltungen für angehende Sportlehrer:innen (u. a. Unterrichtspraktische Studien). Auf dem YouTube-Kanal „Sportunterricht“ zeigt er Spiele und Übungen so, wie sie im echten Unterricht funktionieren.</p><p><b>{"Alle Bände, die App und kostenlose Spiele: " + WEBSITE if WEBSITE else "Kanal: youtube.com/@Sportunterricht2022"}</b></p></div></div></section>'''


def fusszeile(b, probe=False):
    t = f'{b["titel"]} · Sportunterricht Band {b["band"]}'
    if b['slug'] == 'gratis':
        return '5 Spiele, die immer funktionieren · Gratis-Starterpaket der Praxis-Reihe Sportunterricht'
    if probe:
        return t + '   ·   Leseprobe'
    return t if FINAL else t + '   ·   ENTWURF – vor Verkauf prüfen'


def schreibe(b, body, name, probe=False):
    OUT.mkdir(exist_ok=True)
    doc = f'<html lang="de"><head><meta charset="utf-8"><title>{html.escape(b["titel"])}</title><meta name="author" content="David Jungreithmayr · Sportunterricht"><meta name="keywords" content="Sportunterricht, Bewegung und Sport, Spiele, Schule"><style>{css(b["farbe"], fusszeile(b, probe))}</style></head><body>{body}</body></html>'
    pfad = OUT / name
    HTML(string=doc, base_url=str(EBOOKS)).write_pdf(pfad)
    return pfad


def leseprobe_ende(b, punkte):
    wo = f'Erhältlich auf <b>{WEBSITE}</b>' if WEBSITE else 'Den Link zum Shop findest du in der Beschreibung jedes Videos auf <b>youtube.com/@Sportunterricht2022</b>'
    li = ''.join(f'<li>{html.escape(p_)}</li>' for p_ in punkte)
    return f'''<section class="text-seite" style="page-break-before:always"><h1>Das war die Leseprobe</h1>
<p>Hat dir gefallen, was du gesehen hast? Im vollständigen Band {b["band"]} bekommst du:</p><ul>{li}</ul>
<p style="font-family:Barlow;font-size:22pt;font-weight:700;color:{b["farbe"]};margin:6mm 0 1mm">{b["preis"]} · sofort als PDF</p>
<p>Oder alle {BAENDE_WORT} Bände im Komplettpaket für {KOMPLETT["preis"]} statt {KOMPLETT["statt"]} – mit dem Jahresplaner als Bonus.</p><p>{wo}</p></section>'''


# ---------------- Bücher mit Einträgen ----------------
def baue_buch(b):
    buch = parse((EBOOKS / b['slug'] / 'inhalt.md').read_text())
    je_kap = [[eintrag_aus_excel(e) for e in k['eintraege']] for k in buch['kapitel']]
    alle = [e for k in je_kap for e in k]
    neu = any(ist_lang(e) for e in alle)
    seiten = 'Jeder Eintrag beginnt auf einer eigenen Seite' if neu else 'Jeder Eintrag steht vollständig auf einer Seite'
    einleitung = f'''<section class="text-seite" id="einleitung"><h1>Bevor du loslegst</h1>{md(buch["einleitung"])}
<div class="legende"><b>So ist jeder Eintrag aufgebaut:</b> Schulstufe (1–4 Volksschule, 5–8 Unterstufe/Mittelschule, 9–13 Oberstufe) · Gruppengröße · Material · realistische Spielzeit in der Stunde (nicht die Videolänge) · passender Stundenteil · Intensität · Lernziel · Beschreibung · Ansagetext zum Vorlesen · Varianten zum Leichter- und Schwerermachen · Sicherheitshinweis. Der QR-Code führt direkt zur Video-Demo. <b>{seiten}</b> – du kannst ihn einzeln ausdrucken und mit in die Halle nehmen.</div></section>'''
    teile = []
    for i, k in enumerate(buch['kapitel']):
        teile.append(kapitel_kopf(i, k, len(je_kap[i]), b['farbe']))
        if k['text']:
            teile.append(f'<div class="ktext">{md(k["text"])}</div>')
        teile += [karte(e, neu) for e in je_kap[i]]
    body = cover(b, len(alle), stufen=stufen_bereich(alle)) + impressum(b) + toc(buch, je_kap, zusatz=[('anhang', 'Übersicht zum Nachschlagen'), ('reihe', 'Die Praxis-Reihe & der Autor')]) + einleitung + ''.join(teile) + anhang(alle) + reihe_seite(b['slug'])
    voll = schreibe(b, body, f'{b["slug"]}.pdf')

    # Leseprobe: Einleitung, Inhalt, die ersten 3 Einträge
    erste = je_kap[0][:3]
    probe_body = cover(b, len(alle), probe=True, stufen=stufen_bereich(alle)) + toc(buch, je_kap, mit_seitenzahlen=False) + einleitung + kapitel_kopf(0, buch['kapitel'][0], len(je_kap[0]), b['farbe']) + ''.join(karte(e, neu) for e in erste)
    jede = 'jedes' if b['einheit'] == 'Spiele' else 'jede'
    punkte = [f'{len(alle)} {b["einheit"]} in {len(buch["kapitel"])} Kapiteln, {jede} mit Video-Demo per QR-Code',
              'Ansagetext zum Vorlesen, Varianten zum Leichter- und Schwerermachen, Sicherheitshinweise',
              'Übersichtstabellen nach Stundenteil und „ganz ohne Material“']
    probe_body += leseprobe_ende(b, punkte) + reihe_seite(b['slug'])
    probe = schreibe(b, probe_body, f'{b["slug"]}-leseprobe.pdf', probe=True)
    return voll, probe, len(alle), alle


# ---------------- Praktikums-Kit ----------------
HALLE = {'klein': 'Kleine Halle', 'halb': 'Halbe Halle', 'ganz': 'Ganze Halle'}


def baue_praktikum(b, fundorte):
    leit = parse((EBOOKS / 'praktikum' / 'leitfaden.md').read_text())
    stunden = json.loads((EBOOKS / 'praktikum' / 'stunden.json').read_text())
    teile = []
    for i, k in enumerate(leit['kapitel']):
        teile.append(kapitel_kopf(i, k, 0, b['farbe']))
        teile.append(f'<div class="ktext">{md(k["text"])}</div>')
    n0 = len(leit['kapitel'])
    teile.append(f'<section class="kapitel" id="k{n0}"><div class="knr">Teil 2</div><h2>12 fertige Stundenbilder</h2><div class="kmeta">6 × Volksschule, 6 × Sekundarstufe I · je 50 Minuten (45 Minuten Bewegungszeit)</div></section>')
    teile.append('<div class="ktext"><p>Jedes Stundenbild besteht aus einer Seite Verlaufsplanung im Tabellenformat und einer Seite mit dem Ablauf im Detail. Alle Spiele stammen aus der Praxis-Reihe; der Verweis „Band/Nr.“ führt zur ausführlichen Anleitung, der QR-Code zur Video-Demo. Lehrplanbezüge sind Vorschläge – bitte mit dem geltenden Lehrplan und den Vorgaben deiner Betreuungslehrkraft abgleichen.</p></div>')
    for s in stunden:
        material = sorted({x['material'] for x in s['bloecke'] if x['material'] and not re.match(r'^\s*keine', x['material'], re.I)})
        zeilen = []
        for x in s['bloecke']:
            ref = fundorte.get(x['videoId'])
            refs = f'<br><small>Band {ref[0]}, Nr. {ref[1]}</small>' if ref else ''
            zeilen.append(f'<tr class="p-{x["phase"]}"><td class="z">{x["von"]}–{x["bis"]}\'</td><td><b>{x["phase"]}</b></td><td><b>{html.escape(x["titel"])}</b>{refs}</td>'
                          f'<td>{html.escape(x["organisation"] or "")}<br><small>{html.escape(x["sozialform"] or "")}</small></td><td>{html.escape(x["material"] or "–")}</td>'
                          f'<td>{html.escape(x["lernziel"] or "")}<br><small>{html.escape(x["lehrplan"] or "")}</small></td><td>{html.escape(x["differenzierung"] or "")}</td><td>{html.escape(x["sicherheit"] or "")}</td></tr>')
        teile.append(f'''<section class="sb" id="s{s["nr"]}"><div class="sb-kopf"><div><div class="knr" style="color:{b["farbe"]};font-family:Barlow;font-weight:700;letter-spacing:.15em">STUNDENBILD {s["nr"]} · {s["stufe"]}. SCHULSTUFE</div><h2>{html.escape(s["titel"])}</h2></div></div>
<div class="sb-daten"><div><span>Stundenziel</span>{html.escape(s["ziel"])}</div><div><span>Schulstufe</span>{s["stufe"]}. ({"Volksschule" if s["stufe"] <= 4 else "Sekundarstufe I"})</div><div><span>Gruppe</span>{s["gruppe"]} Kinder · {HALLE[s["halle"]]}</div><div><span>Dauer</span>50 min (45 min Bewegung)</div><div><span>Material</span>{html.escape(", ".join(material) or "keines")}</div></div>
<table><thead><tr><th>Zeit</th><th>Phase</th><th>Inhalt</th><th>Organisation / Sozialform</th><th>Material</th><th>Lernziel / Lehrplanbezug</th><th>Differenzierung</th><th>Sicherheit</th></tr></thead><tbody>{"".join(zeilen)}</tbody></table></section>''')
        bl = []
        for x in s['bloecke']:
            if not x['videoId']:
                continue
            ref = fundorte.get(x['videoId'])
            q = f'<img src="{qr("https://www.youtube.com/watch?v=" + x["videoId"])}" style="width:16mm;float:right;margin-left:2mm"/>'
            bl.append(f'<div class="blk">{q}<h4>{x["von"]}–{x["bis"]}\' · {html.escape(x["titel"])}</h4><p>{html.escape(x["beschreibung"])}</p>'
                      f'{"<p class=ansage>„" + html.escape(x["ansage"]) + "“</p>" if x["ansage"] else ""}{"<p><small>Ausführlich: Band " + str(ref[0]) + ", Nr. " + str(ref[1]) + "</small></p>" if ref else ""}</div>')
        bl.append('<div class="blk reflexion"><h4>Reflexion nach der Stunde</h4><p>Wurde das Stundenziel erreicht? Woran habe ich es gesehen?</p><p>Wie viel echte Bewegungszeit hatten die Kinder? Wo gab es Wartezeiten?</p><p>Was ändere ich beim nächsten Mal?</p></div>')
        teile.append(f'<section class="sb-detail"><h3 style="column-span:all;font-size:13pt;margin-bottom:3mm">Ablauf im Detail – Stundenbild {s["nr"]}</h3>{"".join(bl)}</section>')
    hat_einl = bool(leit['einleitung'].strip())
    toc_z = ['<section class="inhalt"><h1>Inhalt</h1><ol class="toc">'] + (['<li class="kap"><a href="#einleitung">Bevor du loslegst</a></li>'] if hat_einl else [])
    toc_z += [f'<li class="kap"><a href="#k{i}">{html.escape(k["titel"])}</a></li>' for i, k in enumerate(leit['kapitel'])]
    toc_z.append(f'<li class="kap"><a href="#k{n0}">Teil 2 – 12 fertige Stundenbilder</a></li>')
    toc_z += [f'<li><span class="nr">{s["nr"]}</span><a href="#s{s["nr"]}">{html.escape(s["titel"])} ({s["stufe"]}. Stufe)</a></li>' for s in stunden]
    toc_z.append('</ol></section>')
    einleitung = f'<section class="text-seite" id="einleitung"><h1>Bevor du loslegst</h1>{md(leit["einleitung"])}</section>' if hat_einl else ''
    body = cover(b, 12) + impressum(b) + ''.join(toc_z) + einleitung + ''.join(teile) + reihe_seite(b['slug'])
    voll = schreibe(b, body, 'praktikum.pdf')
    # Leseprobe: Leitfaden-Kapitel 1 + Stundenbild 1
    probe_body = cover(b, 12, probe=True) + ''.join(toc_z).replace('class="inhalt"', 'class="inhalt ohne-seiten"') + einleitung + teile[0] + teile[1]
    i1 = next(i for i, t in enumerate(teile) if 'id="s1"' in t)
    probe_body += teile[i1] + teile[i1 + 1] + leseprobe_ende(b, ['7 Kapitel Leitfaden: Aufbau, Lernziele, Lehrplanbezug, Organisation, Differenzierung, Reflexion, Checkliste',
                                                               '12 fertige Stundenbilder (6 × Volksschule, 6 × Sekundarstufe I) mit Verlaufsplanung und Ablauf im Detail',
                                                               'Jedes Spiel mit Verweis auf die ausführliche Anleitung und QR-Code zur Video-Demo']) + reihe_seite(b['slug'])
    probe = schreibe(b, probe_body, 'praktikum-leseprobe.pdf', probe=True)
    return voll, probe


# ---------------- Bonus: Jahresplaner ----------------
# Redaktionell zusammengestellt: (Monat, Schwerpunkt, Tipp, Volksschule [(Band, Nr)], Sekundarstufe [(Band, Nr)])
MONATE = [
    ('September', 'Ankommen & Kooperation', 'Regeln, Signale und Rituale einführen – die ersten Wochen entscheiden über das Klassenklima.',
     [(1, 9), (1, 13), (1, 28), (1, 50)], [(1, 10), (1, 14), (1, 16), (1, 11)]),
    ('Oktober', 'Fangen, Laufen, Ausdauer', 'Ausdauer spielerisch aufbauen – und draußen laufen, solange das Wetter hält.',
     [(2, 5), (2, 3), (2, 12), (1, 18)], [(2, 16), (4, 9), (4, 8), (2, 19)]),
    ('November', 'Turnen am Boden', 'Rollen, Stützen, Balancieren – Mattenaufbau und Hilfestellung früh als Routine einüben.',
     [(3, 23), (3, 1), (3, 15), (4, 19)], [(3, 3), (3, 4), (3, 5), (3, 6)]),
    ('Dezember', 'Springen & Stützen', 'Sprünge mit sicherer Landung; Stationsbetrieb mit festen Aufbauteams spart Zeit.',
     [(3, 11), (3, 8), (3, 25), (2, 33)], [(3, 9), (3, 10), (3, 13), (3, 16)]),
    ('Jänner', 'Fitness & Zirkel', 'Kraft und Körperspannung mit dem eigenen Körpergewicht – Zirkel machen Fortschritt sichtbar.',
     [(4, 4), (4, 15), (4, 22), (4, 28)], [(4, 1), (4, 23), (4, 24), (4, 25)]),
    ('Februar', 'Kleine Spiele mit Ball', 'Zuspielen, Freilaufen, Fairplay – Ballgefühl vor Regelwerk.',
     [(2, 29), (1, 39), (1, 30), (1, 37)], [(1, 35), (1, 32), (1, 33), (1, 38)]),
    ('März', 'Akrobatik & Gestalten', 'Partner- und Gruppenakrobatik mit einer kleinen Präsentation als Abschluss.',
     [(3, 28), (3, 27), (1, 12), (1, 27)], [(3, 27), (3, 29), (3, 30), (1, 17)]),
    ('April', 'Leichtathletik spielerisch', 'Laufen, Springen, Werfen in Spielformen – Messen macht den eigenen Fortschritt sichtbar.',
     [(4, 29), (4, 30), (4, 34), (2, 11)], [(4, 31), (4, 32), (4, 33), (4, 35)]),
    ('Mai', 'Koordination & Rückschlagspiele', 'Reaktion, Gleichgewicht, Ballgefühl – viele Stationen, wenig Warten.',
     [(4, 16), (4, 20), (4, 21), (1, 41)], [(4, 17), (4, 18), (1, 45), (1, 31)]),
    ('Juni', 'Spielfest & Teamchallenges', 'Große Spiele, Teamwettbewerbe und ein gemeinsamer Abschluss des Schuljahres.',
     [(1, 29), (1, 22), (1, 19), (1, 21)], [(1, 26), (1, 29), (1, 48), (1, 24)]),
]


def baue_jahresplan(fundorte, alle_eintraege):
    nach = {(band, e['nr']): e for e, band in alle_eintraege}
    probleme = []

    def ideen(liste, sek, monat):
        z = []
        for band, nr in liste:
            e = nach.get((band, nr))
            if not e:
                probleme.append(f'{monat}: B{band}/{nr} fehlt')
                continue
            v = VIDEOS.get(e['id']) or {}
            von, bis = int(v.get('Stufe_von') or 1), int(v.get('Stufe_bis') or 13)
            if (not sek and von > 4) or (sek and (bis < 5 or von > 8)):
                probleme.append(f'{monat}: {e["name"]} ({von}–{bis}) passt nicht zu {"Sek I" if sek else "VS"}')
            z.append(f'<li><b>B{band}/{nr}</b> {html.escape(e["name"])} <small>({von}.–{bis}.)</small></li>')
        return '<ul class="jp">' + ''.join(z) + '</ul>'
    zeilen = []
    for monat, thema, tipp, vs, sek in MONATE:
        zeilen.append(f'<tr><td class="mon">{monat}</td><td><b>{thema}</b><br><small>{tipp}</small></td><td>{ideen(vs, False, monat)}</td><td>{ideen(sek, True, monat)}</td><td class="notiz"></td></tr>')
    for p in probleme:
        print('⚠ Jahresplaner:', p)
    kopf = '<table class="jpt"><thead><tr><th style="width:21mm">Monat</th><th style="width:62mm">Schwerpunkt</th><th>Volksschule (1.–4.)</th><th>Sekundarstufe I (5.–8.)</th><th style="width:48mm">Meine Notizen</th></tr></thead><tbody>'
    leer = ''.join(f'<tr><td class="mon">{m[0]}</td><td></td><td></td><td></td><td></td></tr>' for m in MONATE)
    b = {'titel': 'Jahresplaner Bewegung und Sport'}
    body = f"""<section class="cover" style="background:#13211b">{FELD_SVG}<div class="reihe">Sportunterricht · Die Praxis-Reihe</div><div class="band">BONUS ZUM KOMPLETTPAKET</div>
<h1>Jahresplaner Bewegung und Sport</h1><p class="unter">Zehn Monatsschwerpunkte von September bis Juni – mit je vier passenden Ideen für Volksschule und Sekundarstufe I aus allen Bänden der Reihe.</p>
<div class="autor">David Jungreithmayr · AHS-Sportlehrer in Wien · YouTube: @Sportunterricht2022</div></section>
<section class="text-seite"><h1>So nutzt du den Plan</h1>
<p>Der Plan ist ein Vorschlag für ein ausgewogenes Schuljahr: Jeder Monat hat einen Schwerpunkt, dazu je vier Ideen für die Volksschule und die Sekundarstufe I. Der Verweis „B3/5“ bedeutet <b>Band 3, Eintrag Nr. 5</b>; in Klammern steht der Schulstufenbereich des Eintrags.</p>
<p>Ein Schwerpunkt heißt nicht, dass der ganze Monat nur daraus besteht. Bewährt hat sich: <b>ein Schwerpunkt im Hauptteil</b>, dazu wechselnde Aufwärm- und Abschlussspiele aus Band 1 und 2. So bleibt Raum für Wünsche der Klasse, Schulveranstaltungen und Wetter.</p>
<p>Am Ende findest du eine <b>leere Vorlage</b> für deinen eigenen Jahresplan – zum Ausdrucken für jede Klasse. Die App schlägt dir zu jedem Schwerpunkt auf Knopfdruck weitere passende Spiele vor.</p>
<h2>Faustregeln für das Schuljahr</h2>
<ul><li><b>Turnen vor Weihnachten:</b> Die Halle ist frei von Außenaktivitäten, die Klasse kennt die Regeln – ideal für Geräte.</li>
<li><b>Leichtathletik ab April:</b> Draußen messen, laufen, werfen; bei Schlechtwetter die Spielformen aus Band 4 in der Halle.</li>
<li><b>Jeder Monat mit einem sichtbaren Ergebnis:</b> eine kleine Präsentation, ein Rekord, ein Turnier – das motiviert mehr als jede Note.</li>
<li><b>Kooperation zu Schuljahresbeginn und vor den Ferien:</b> Sie tut dem Klassenklima am meisten gut, wenn sich die Gruppe neu findet oder verabschiedet.</li></ul></section>
<section class="sb" style="page-break-before:always"><div class="sb-kopf"><div><div class="knr">SEPTEMBER BIS JÄNNER</div><h2>Jahresplan – erstes Semester</h2></div></div>{kopf}{''.join(zeilen[:5])}</tbody></table></section>
<section class="sb"><div class="sb-kopf"><div><div class="knr">FEBRUAR BIS JUNI</div><h2>Jahresplan – zweites Semester</h2></div></div>{kopf}{''.join(zeilen[5:])}</tbody></table></section>
<section class="sb"><div class="sb-kopf"><div><div class="knr">ZUM AUSFÜLLEN</div><h2>Mein Jahresplan · Klasse: ____________ · Schuljahr: ____________</h2></div></div>{kopf.replace('<table class="jpt">', '<table class="jpt leer">').replace('>Schwerpunkt<', '>Thema<').replace('Volksschule (1.–4.)', 'Hauptteil-Ideen').replace('Sekundarstufe I (5.–8.)', 'Aufwärmen & Abschluss')}{leer}</tbody></table></section>""" + reihe_seite('')
    extra = """
.knr { font-family: 'Barlow'; font-weight: 700; letter-spacing: .15em; font-size: 8.5pt; color: #167a4b; }
.jpt { font-size: 8pt; }
.jpt td { padding: 1.8mm 2mm; }
.jpt td.mon { font-family: 'Barlow'; font-weight: 700; font-size: 12pt; }
.jpt td.notiz { border-left: 0.6pt dashed #d2ddd7; }
.jpt tbody tr { height: 30mm; }
.jpt.leer tbody tr { height: 15.5mm; }
ul.jp { list-style: none; padding: 0; margin: 0; }
ul.jp li { margin-bottom: 0.7mm; }
ul.jp small { color: #66776f; }
"""
    doc = f'<html lang="de"><head><meta charset="utf-8"><title>Jahresplaner</title><style>{css("#13211b", "Jahresplaner · Sportunterricht Praxis-Reihe · Bonus zum Komplettpaket")}{extra}</style></head><body>{body}</body></html>'
    HTML(string=doc, base_url=str(EBOOKS)).write_pdf(OUT / 'jahresplaner.pdf')


# ---------------- Gratis-Starterpaket (Lead-Magnet für die Website) ----------------
GRATIS = [(1, 5), (1, 8), (1, 11), (1, 25), (1, 24)]


def baue_gratis(alle_eintraege):
    nach = {(band, e['nr']): e for e, band in alle_eintraege}
    b = {'slug': 'gratis', 'band': 1, 'titel': '5 Spiele, die immer funktionieren', 'unter': 'Das kostenlose Starterpaket: ohne Material, für jede Halle, mit Video-Demo und Ansagetext',
         'farbe': '#167a4b', 'einheit': 'Spiele', 'preis': 'kostenlos'}
    eintraege = [nach[k] for k in GRATIS]
    karten = []
    for i, e in enumerate(eintraege, 1):
        karten.append(karte({**e, 'nr': i}))
    herkunft = ', '.join(f'„{e["name"]}“ (Band 1, Nr. {n})' for (_, n), e in zip(GRATIS, eintraege))
    body = f'''<section class="cover" style="background:#167a4b">{FELD_SVG}<div class="reihe">Sportunterricht · Gratis-Starterpaket</div><div class="band">KOSTENLOS</div>
<h1>5 Spiele, die immer funktionieren</h1><p class="unter">Ohne Material, für jede Halle, von der Volksschule bis zur Oberstufe – mit Video-Demo und Ansagetext zum Vorlesen.</p>
<div class="zahlen"><div><b>5</b>Spiele</div><div><b>0</b>Material</div><div><b>5</b>Video-Demos</div></div>
<div class="autor">David Jungreithmayr · AHS-Sportlehrer in Wien · YouTube: @Sportunterricht2022</div></section>
<section class="text-seite"><h1>Schön, dass du da bist</h1>
<p>Diese fünf Spiele sind meine Joker für Tage, an denen nichts nach Plan läuft: Die Halle ist geteilt, der Geräteraum zu, die Klasse unruhig. Sie brauchen kein Material, sind in einer Minute erklärt und funktionieren mit einer ganzen Klasse von 15 bis 30 Kindern.</p>
<p>Jede Karte hat denselben Aufbau wie in der Praxis-Reihe: Schulstufe, Gruppe, Dauer, Lernziel, Beschreibung, ein Ansagetext zum Vorlesen und Varianten zum Leichter- und Schwerermachen. Der QR-Code führt direkt zur Video-Demo.</p>
<h2>So setzt du sie ein</h2><ul><li><b>Aufwärmen:</b> Schwarz-Weiß, Countdown, Magneto</li><li><b>Kooperation im Hauptteil:</b> Gordischer Knoten</li><li><b>Ruhiger Abschluss:</b> Sortieren (oder noch einmal der Gordische Knoten)</li></ul>
<p class="legende">Alle fünf Spiele stammen aus Band 1 „50 kreative &amp; kooperative Spiele mit wenig Material“: {html.escape(herkunft)}.</p></section>
<section class="kapitel" id="k0">{FELD_KLEIN}<div class="knr">Gratis-Starterpaket</div><h2>Die 5 Spiele</h2></section>{''.join(karten)}''' + reihe_seite('')
    schreibe(b, body, 'gratis-5-spiele.pdf', probe=True)


def cover_bild(pdf, name):
    subprocess.run(['pdftoppm', '-png', '-r', '70', '-f', '1', '-l', '1', '-singlefile', str(pdf), str(OUT / name)], check=True)


def main():
    OUT.mkdir(exist_ok=True)
    fundorte = {}
    alle_eintraege = []
    for b in REIHE:
        if b['slug'] in ('praktikum', 'alle-dabei'):
            continue
        voll, probe, n, alle = baue_buch(b)
        for e in alle:
            fundorte.setdefault(e['id'], (b['band'], e['nr']))
            alle_eintraege.append((e, b['band']))
        cover_bild(voll, f'{b["slug"]}-cover')
        print(f'✓ Band {b["band"]} {b["kurz"]}: {n} Einträge → {voll.name}, {probe.name}')
    p = REIHE[4]
    voll, probe = baue_praktikum(p, fundorte)
    cover_bild(voll, 'praktikum-cover')
    print(f'✓ Band 5 Praktikums-Kit → {voll.name}, {probe.name}')
    from alle_dabei import baue_alle_dabei
    voll, probe, info = baue_alle_dabei(REIHE[5], fundorte)
    cover_bild(voll, 'alle-dabei-cover')
    print(f'✓ Band 6 Alle dabei: {info} → {voll.name}, {probe.name}')
    baue_jahresplan(fundorte, alle_eintraege)
    baue_gratis(alle_eintraege)
    print('✓ Gratis-Starterpaket')
    cover_bild(OUT / 'jahresplaner.pdf', 'jahresplaner-cover')
    print('✓ Bonus Jahresplaner')
    (OUT / 'fundorte.json').write_text(json.dumps({k: list(v) for k, v in fundorte.items()}))


if __name__ == '__main__':
    main()
